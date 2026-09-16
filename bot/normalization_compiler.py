from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import UTC, datetime
import json
from pathlib import Path
import re
from typing import Any
import unicodedata


RUNTIME_VERSION = 1

RELEVANT_SHEETS = (
    "All_aliases",
    "Categories",
    "Conflicts",
    "FoodService_aliases",
    "FoodService_merchants",
    "FoodService_terms",
    "Merchant_aliases",
    "Merchants_177",
    "Product_aliases",
    "Products_341",
    "Regex_rules",
    "Stopwords_noise",
    "Subscription_aliases",
    "Subscriptions",
)

_SPACE_RE = re.compile(r"\s+")
_NON_WORD_RE = re.compile(r"[^0-9a-z\u0400-\u052f']+")

_CATEGORY_SLUG_BY_NAME = {
    "\u043f\u0440\u043e\u0434\u0443\u043a\u0442\u0438": "groceries",
    "\u0434\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u043f\u0440\u043e\u0434\u0443\u043a\u0442\u0456\u0432": "groceries",
    "\u043a\u0430\u0444\u0435 \u0442\u0430 \u0440\u0435\u0441\u0442\u043e\u0440\u0430\u043d\u0438": "cafes_restaurants_delivery",
    "\u0434\u043e\u0441\u0442\u0430\u0432\u043a\u0430 \u0457\u0436\u0456": "cafes_restaurants_delivery",
    "\u0440\u0435\u0441\u0442\u043e\u0440\u0430\u043d\u0438": "cafes_restaurants_delivery",
    "\u043a\u0430\u0432\u0430 \u0442\u0430 \u043a\u0430\u0432'\u044f\u0440\u043d\u0456": "cafes_restaurants_delivery",
    "\u043a\u0430\u0432\u0430 \u0442\u0430 \u043a\u0430\u0432\u2019\u044f\u0440\u043d\u0456": "cafes_restaurants_delivery",
    "\u043f\u0435\u043a\u0430\u0440\u043d\u044f \u0432\u0438\u043f\u0456\u0447\u043a\u0430": "cafes_restaurants_delivery",
    "\u0431\u0430\u0440\u0438 \u043f\u0430\u0431\u0438": "cafes_restaurants_delivery",
    "\u0430\u0437\u0456\u0439\u0441\u044c\u043a\u0430 \u043a\u0443\u0445\u043d\u044f": "cafes_restaurants_delivery",
    "\u043f\u0456\u0446\u0430": "cafes_restaurants_delivery",
    "\u0441\u0443\u0448\u0456": "cafes_restaurants_delivery",
    "\u0444\u0430\u0441\u0442\u0444\u0443\u0434": "cafes_restaurants_delivery",
    "\u0448\u0430\u0443\u0440\u043c\u0430 \u0434\u043e\u043d\u0435\u0440 \u043a\u0435\u0431\u0430\u0431": "cafes_restaurants_delivery",
    "\u043e\u0440\u0435\u043d\u0434\u0430 \u0436\u0438\u0442\u043b\u0430": "rent",
    "\u043a\u043e\u043c\u0443\u043d\u0430\u043b\u044c\u043d\u0456 \u043f\u043b\u0430\u0442\u0435\u0436\u0456": "utilities",
    "\u0434\u0456\u043c \u0442\u0430 \u0440\u0435\u043c\u043e\u043d\u0442": "home_household",
    "\u0434\u0456\u043c \u0442\u0430 \u043f\u043e\u0431\u0443\u0442": "home_household",
    "\u043f\u043e\u0431\u0443\u0442\u043e\u0432\u0430 \u0445\u0456\u043c\u0456\u044f": "home_household",
    "\u0442\u0440\u0430\u043d\u0441\u043f\u043e\u0440\u0442": "transport",
    "\u0430\u0432\u0442\u043e": "transport",
    "\u043f\u0430\u043b\u044c\u043d\u0435": "transport",
    "\u0442\u0430\u043a\u0441\u0456": "taxi",
    "\u0437\u0432'\u044f\u0437\u043e\u043a \u0442\u0430 \u0456\u043d\u0442\u0435\u0440\u043d\u0435\u0442": "telecom_internet",
    "\u0437\u0432\u02bc\u044f\u0437\u043e\u043a \u0442\u0430 \u0456\u043d\u0442\u0435\u0440\u043d\u0435\u0442": "telecom_internet",
    "\u0430\u043f\u0442\u0435\u043a\u0430 \u0442\u0430 \u043b\u0456\u043a\u0438": "health",
    "\u043c\u0435\u0434\u0438\u0446\u0438\u043d\u0430": "health",
    "\u0437\u0434\u043e\u0440\u043e\u0432'\u044f": "health",
    "\u0437\u0434\u043e\u0440\u043e\u0432\u2019\u044f": "health",
    "\u043e\u0434\u044f\u0433 \u0442\u0430 \u0432\u0437\u0443\u0442\u0442\u044f": "clothes_shoes",
    "\u043a\u0440\u0430\u0441\u0430 \u0442\u0430 \u0434\u043e\u0433\u043b\u044f\u0434": "beauty_care",
    "\u043f\u0456\u0434\u043f\u0438\u0441\u043a\u0438 \u0442\u0430 \u0441\u0435\u0440\u0432\u0456\u0441\u0438": "subscriptions_services",
    "\u0456\u0433\u0440\u0438": "entertainment_leisure",
    "\u0440\u043e\u0437\u0432\u0430\u0433\u0438": "entertainment_leisure",
    "\u043a\u0456\u043d\u043e \u0442\u0430 \u0440\u043e\u0437\u0432\u0430\u0433\u0438": "entertainment_leisure",
    "\u0441\u043f\u043e\u0440\u0442": "entertainment_leisure",
    "\u043f\u043e\u0434\u043e\u0440\u043e\u0436\u0456": "entertainment_leisure",
    "\u043f\u043e\u0434\u0430\u0440\u0443\u043d\u043a\u0438": "gifts_donations",
    "\u0431\u043b\u0430\u0433\u043e\u0434\u0456\u0439\u043d\u0456\u0441\u0442\u044c": "gifts_donations",
    "\u0456\u043d\u0448\u0435": "other",
}

_MERCHANT_TYPE_SLUG_OVERRIDES = {
    "taxi": "taxi",
    "food_delivery_app": "cafes_restaurants_delivery",
    "pizza": "cafes_restaurants_delivery",
    "restaurant": "cafes_restaurants_delivery",
    "restaurant_group": "cafes_restaurants_delivery",
    "restaurant_bar": "cafes_restaurants_delivery",
    "restaurant_cafe": "cafes_restaurants_delivery",
    "italian_restaurant": "cafes_restaurants_delivery",
    "asian_restaurant": "cafes_restaurants_delivery",
    "seafood_restaurant": "cafes_restaurants_delivery",
    "kebab_restaurant": "cafes_restaurants_delivery",
    "coffee": "cafes_restaurants_delivery",
    "cafe": "cafes_restaurants_delivery",
    "dessert_cafe": "cafes_restaurants_delivery",
    "bakery_cafe": "cafes_restaurants_delivery",
    "fast_food": "cafes_restaurants_delivery",
    "fast_casual": "cafes_restaurants_delivery",
    "turkish_fast_casual": "cafes_restaurants_delivery",
    "sandwich_fast_casual": "cafes_restaurants_delivery",
    "sandwich_fast_food": "cafes_restaurants_delivery",
    "healthy_food": "cafes_restaurants_delivery",
    "bar": "cafes_restaurants_delivery",
    "bar_cafe": "cafes_restaurants_delivery",
    "beer_restaurant": "cafes_restaurants_delivery",
    "grocery": "groceries",
    "grocery_discount": "groceries",
    "grocery_cash_carry": "groceries",
    "grocery_delivery": "groceries",
    "convenience": "groceries",
    "food_specialty": "groceries",
    "food_wine": "groceries",
    "alcohol_snacks": "groceries",
    "pharmacy": "health",
    "pharmacy_online": "health",
    "marketplace_health": "health",
    "beauty": "beauty_care",
    "beauty_online": "beauty_care",
    "drogerie": "beauty_care",
}

_ROUTING_CARRIER_TYPES = {"food_delivery_app"}


def normalize_text(value: Any) -> str:
    raw = unicodedata.normalize("NFKC", str(value or ""))
    raw = raw.casefold()
    raw = raw.replace("\u2019", "'").replace("\u02bc", "'").replace("`", "'")
    raw = _NON_WORD_RE.sub(" ", raw)
    return _SPACE_RE.sub(" ", raw).strip()


def _normalize_type_value(value: Any) -> str:
    raw = str(value or "").strip().casefold()
    raw = raw.replace("-", "_").replace(" ", "_")
    return raw


def _slug_for_category_name(value: Any) -> str | None:
    return _CATEGORY_SLUG_BY_NAME.get(normalize_text(value))


def _safe_int(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _sheet_rows_by_name(workbook_path: Path) -> dict[str, list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - runtime dependency path
        raise RuntimeError("openpyxl is required to compile normalization dictionaries") from exc

    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    rows_by_name: dict[str, list[dict[str, Any]]] = {}
    for sheet_name in RELEVANT_SHEETS:
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        iterator = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration:
            rows_by_name[sheet_name] = []
            continue
        headers = [str(cell or "").strip() for cell in header_row]
        records: list[dict[str, Any]] = []
        for row in iterator:
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            record: dict[str, Any] = {}
            for index, header in enumerate(headers):
                if not header:
                    continue
                record[header] = row[index] if index < len(row) else None
            records.append(record)
        rows_by_name[sheet_name] = records
    return rows_by_name


def _new_report(source_path: str | None) -> dict[str, Any]:
    return {
        "runtime_version": RUNTIME_VERSION,
        "source_path": source_path or "",
        "generated_at": datetime.now(UTC).isoformat(),
        "counts": {},
        "unmapped_categories": [],
        "duplicate_alias_resolutions": [],
        "invalid_regex_rules": [],
        "skipped_short_aliases": [],
        "skipped_conflict_aliases": [],
    }


def _report_count(report: dict[str, Any], key: str, amount: int = 1) -> None:
    counts = report["counts"]
    counts[key] = int(counts.get(key, 0)) + amount


def _track_unmapped(unmapped: dict[str, dict[str, Any]], category_name: Any, source: str) -> None:
    name = str(category_name or "").strip()
    if not name:
        return
    bucket = unmapped.setdefault(name, {"count": 0, "sources": set()})
    bucket["count"] += 1
    bucket["sources"].add(source)


def _resolve_duplicate(
    existing: dict[str, Any] | None,
    candidate: dict[str, Any],
    *,
    bucket_name: str,
    report: dict[str, Any],
) -> dict[str, Any]:
    if existing is None:
        return candidate

    current_key = (
        _safe_int(existing.get("priority"), 0),
        len(str(existing.get("alias", ""))),
        str(existing.get("matched_slug", "")),
        str(existing.get("normalize_to", "")),
    )
    candidate_key = (
        _safe_int(candidate.get("priority"), 0),
        len(str(candidate.get("alias", ""))),
        str(candidate.get("matched_slug", "")),
        str(candidate.get("normalize_to", "")),
    )
    if candidate_key > current_key:
        report["duplicate_alias_resolutions"].append(
            {
                "bucket": bucket_name,
                "alias": candidate.get("alias", ""),
                "kept_slug": candidate.get("matched_slug", ""),
                "replaced_slug": existing.get("matched_slug", ""),
            }
        )
        return candidate
    if candidate_key < current_key:
        report["duplicate_alias_resolutions"].append(
            {
                "bucket": bucket_name,
                "alias": candidate.get("alias", ""),
                "kept_slug": existing.get("matched_slug", ""),
                "discarded_slug": candidate.get("matched_slug", ""),
            }
        )
    return existing


def _register_alias(
    bucket: dict[str, dict[str, Any]],
    alias_raw: Any,
    candidate: dict[str, Any],
    *,
    bucket_name: str,
    report: dict[str, Any],
    allow_conflict_alias: bool,
    conflict_aliases: set[str],
) -> None:
    alias = normalize_text(alias_raw)
    if len(alias) < 3:
        report["skipped_short_aliases"].append({"bucket": bucket_name, "alias": alias, "raw": str(alias_raw or "")})
        return
    if not allow_conflict_alias and alias in conflict_aliases:
        report["skipped_conflict_aliases"].append({"bucket": bucket_name, "alias": alias})
        return
    candidate = dict(candidate)
    candidate["alias"] = alias
    bucket[alias] = _resolve_duplicate(bucket.get(alias), candidate, bucket_name=bucket_name, report=report)


def _build_merchant_meta(
    rows_by_name: dict[str, list[dict[str, Any]]],
) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    meta: dict[str, dict[str, Any]] = {}
    unmapped: dict[str, dict[str, Any]] = {}

    for sheet_name, name_key, type_key, category_key, subcategory_key in (
        ("Merchants_177", "merchant_norm", "merchant_type", "default_category", None),
        ("FoodService_merchants", "merchant_norm", "merchant_type", "default_category", "subcategory"),
    ):
        for row in rows_by_name.get(sheet_name, []):
            merchant_name = normalize_text(row.get(name_key))
            if not merchant_name:
                continue
            merchant_type = _normalize_type_value(row.get(type_key))
            slug = _MERCHANT_TYPE_SLUG_OVERRIDES.get(merchant_type) or _slug_for_category_name(row.get(category_key))
            if slug is None:
                _track_unmapped(unmapped, row.get(category_key), sheet_name)
                continue
            meta[merchant_name] = {
                "normalize_to": str(row.get(name_key) or "").strip(),
                "matched_slug": slug,
                "matched_entity": "merchant",
                "merchant_type": merchant_type,
                "subcategory": str(row.get(subcategory_key) or "").strip() or None,
                "routing_carrier": merchant_type in _ROUTING_CARRIER_TYPES,
            }

    for row in rows_by_name.get("Subscriptions", []):
        service_name = normalize_text(row.get("service_norm"))
        if not service_name:
            continue
        slug = _slug_for_category_name(row.get("default_category"))
        if slug is None:
            _track_unmapped(unmapped, row.get("default_category"), "Subscriptions")
            continue
        meta[service_name] = {
            "normalize_to": str(row.get("service_norm") or "").strip(),
            "matched_slug": slug,
            "matched_entity": "subscription",
            "merchant_type": _normalize_type_value(row.get("service_type")),
            "subcategory": str(row.get("service_type") or "").strip() or None,
            "routing_carrier": False,
        }

    return meta, unmapped


def compile_rows_by_sheet(rows_by_name: dict[str, list[dict[str, Any]]], *, source_path: str | None = None) -> tuple[dict[str, Any], dict[str, Any]]:
    report = _new_report(source_path)
    merchant_meta, unmapped_categories = _build_merchant_meta(rows_by_name)
    conflict_aliases = {
        normalize_text(row.get("alias_raw"))
        for row in rows_by_name.get("Conflicts", [])
        if normalize_text(row.get("alias_raw"))
    }

    exact_aliases: dict[str, dict[str, Any]] = {}
    product_aliases: dict[str, dict[str, Any]] = {}
    category_keywords: dict[str, dict[str, Any]] = {}
    regex_rules: list[dict[str, Any]] = []

    for sheet_name in ("Merchant_aliases", "FoodService_aliases", "Subscription_aliases"):
        for row in rows_by_name.get(sheet_name, []):
            target_key = normalize_text(row.get("normalize_to"))
            meta = merchant_meta.get(target_key)
            slug = meta.get("matched_slug") if meta else _slug_for_category_name(row.get("main_category"))
            if not slug:
                _track_unmapped(unmapped_categories, row.get("main_category"), sheet_name)
                continue
            matched_entity = (meta or {}).get("matched_entity") or normalize_text(row.get("entity_type")) or "merchant"
            candidate = {
                "normalize_to": (meta or {}).get("normalize_to") or str(row.get("normalize_to") or "").strip(),
                "matched_slug": slug,
                "matched_entity": matched_entity,
                "subcategory": (meta or {}).get("subcategory") or str(row.get("subcategory_or_type") or "").strip() or None,
                "priority": _safe_int(row.get("priority"), 0),
                "routing_carrier": bool((meta or {}).get("routing_carrier")),
            }
            _register_alias(
                exact_aliases,
                row.get("alias_raw"),
                candidate,
                bucket_name="exact_aliases",
                report=report,
                allow_conflict_alias=True,
                conflict_aliases=conflict_aliases,
            )

    for sheet_name, alias_key, category_key, normalize_key, subcategory_key in (
        ("Product_aliases", "alias_raw", "main_category", "normalize_to", "subcategory_or_type"),
        ("Categories", "keyword_or_phrase", "normalize_to_category", "normalize_to_category", "category_group"),
        ("All_aliases", "alias_raw", "main_category", "normalize_to", "subcategory_or_type"),
    ):
        for row in rows_by_name.get(sheet_name, []):
            entity_type = normalize_text(row.get("entity_type"))
            if sheet_name == "All_aliases" and entity_type not in {"category", "product"}:
                continue
            slug = _slug_for_category_name(row.get(category_key))
            if slug is None:
                _track_unmapped(unmapped_categories, row.get(category_key), sheet_name)
                continue
            bucket = product_aliases if sheet_name == "Product_aliases" or entity_type == "product" else category_keywords
            bucket_name = "product_aliases" if bucket is product_aliases else "category_keywords"
            candidate = {
                "normalize_to": str(row.get(normalize_key) or "").strip(),
                "matched_slug": slug,
                "matched_entity": entity_type or ("product" if bucket is product_aliases else "category"),
                "subcategory": str(row.get(subcategory_key) or "").strip() or None,
                "priority": _safe_int(row.get("priority"), 0),
                "routing_carrier": False,
            }
            _register_alias(
                bucket,
                row.get(alias_key),
                candidate,
                bucket_name=bucket_name,
                report=report,
                allow_conflict_alias=bucket is product_aliases,
                conflict_aliases=conflict_aliases,
            )

    for index, row in enumerate(rows_by_name.get("Regex_rules", []), start=1):
        pattern = str(row.get("regex_pattern") or "").strip()
        if not pattern:
            continue
        slug = _slug_for_category_name(row.get("normalize_to"))
        if slug is None:
            _track_unmapped(unmapped_categories, row.get("normalize_to"), "Regex_rules")
            continue
        try:
            re.compile(pattern)
        except re.error as exc:
            report["invalid_regex_rules"].append({"pattern": pattern, "error": str(exc)})
            continue
        regex_rules.append(
            {
                "pattern": pattern,
                "matched_slug": slug,
                "matched_entity": normalize_text(row.get("entity_type")) or "category_regex",
                "subcategory": str(row.get("main_category_or_subcategory") or "").strip() or None,
                "priority": _safe_int(index, 0),
            }
        )

    stopwords = sorted(
        {
            normalize_text(row.get("noise_word"))
            for row in rows_by_name.get("Stopwords_noise", [])
            if normalize_text(row.get("noise_word"))
        },
        key=lambda item: (-len(item), item),
    )

    routing_terms = sorted(
        {
            normalize_text(row.get("keyword_or_phrase"))
            for row in rows_by_name.get("FoodService_terms", [])
            if normalize_text(row.get("keyword_or_phrase"))
        },
        key=lambda item: (-len(item), item),
    )

    runtime = {
        "runtime_version": RUNTIME_VERSION,
        "generated_at": datetime.now(UTC).isoformat(),
        "source_path": source_path or "",
        "exact_aliases": sorted(exact_aliases.values(), key=lambda item: (-len(str(item["alias"])), -_safe_int(item.get("priority"), 0), str(item["alias"]))),
        "product_aliases": sorted(product_aliases.values(), key=lambda item: (-len(str(item["alias"])), -_safe_int(item.get("priority"), 0), str(item["alias"]))),
        "category_keywords": sorted(category_keywords.values(), key=lambda item: (-len(str(item["alias"])), -_safe_int(item.get("priority"), 0), str(item["alias"]))),
        "regex_rules": regex_rules,
        "routing": {
            "carrier_aliases": sorted(
                [item["alias"] for item in exact_aliases.values() if item.get("routing_carrier")],
                key=lambda item: (-len(item), item),
            ),
            "food_terms": routing_terms,
        },
        "stopwords": stopwords,
        "conflict_aliases": sorted(conflict_aliases),
    }

    _report_count(report, "merchant_meta", len(merchant_meta))
    _report_count(report, "exact_aliases", len(runtime["exact_aliases"]))
    _report_count(report, "product_aliases", len(runtime["product_aliases"]))
    _report_count(report, "category_keywords", len(runtime["category_keywords"]))
    _report_count(report, "regex_rules", len(runtime["regex_rules"]))
    _report_count(report, "carrier_aliases", len(runtime["routing"]["carrier_aliases"]))
    _report_count(report, "food_terms", len(runtime["routing"]["food_terms"]))
    _report_count(report, "stopwords", len(runtime["stopwords"]))
    report["unmapped_categories"] = [
        {"category": name, "count": payload["count"], "sources": sorted(payload["sources"])}
        for name, payload in sorted(unmapped_categories.items())
    ]

    return runtime, report


def compile_workbook(
    workbook_path: Path,
    *,
    runtime_output_path: Path | None = None,
    report_output_path: Path | None = None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    rows_by_name = _sheet_rows_by_name(workbook_path)
    runtime, report = compile_rows_by_sheet(rows_by_name, source_path=str(workbook_path))

    if runtime_output_path is not None:
        runtime_output_path.write_text(json.dumps(runtime, ensure_ascii=False, indent=2), encoding="utf-8")
    if report_output_path is not None:
        report_output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return runtime, report


def _default_artifact_path(filename: str) -> Path:
    return Path(__file__).with_name(filename)


def main() -> int:
    parser = argparse.ArgumentParser(description="Compile expense normalization workbook into a runtime artifact.")
    parser.add_argument("workbook", type=Path, help="Path to the source .xlsx workbook.")
    parser.add_argument(
        "--runtime-output",
        type=Path,
        default=_default_artifact_path("normalization_runtime.json"),
        help="Path to the generated runtime artifact JSON.",
    )
    parser.add_argument(
        "--report-output",
        type=Path,
        default=_default_artifact_path("normalization_validation_report.json"),
        help="Path to the generated validation report JSON.",
    )
    args = parser.parse_args()

    runtime_output = Path(args.runtime_output)
    report_output = Path(args.report_output)
    runtime_output.parent.mkdir(parents=True, exist_ok=True)
    report_output.parent.mkdir(parents=True, exist_ok=True)

    compile_workbook(
        Path(args.workbook),
        runtime_output_path=runtime_output,
        report_output_path=report_output,
    )
    return 0


if __name__ == "__main__":  # pragma: no cover - CLI entry point
    raise SystemExit(main())
